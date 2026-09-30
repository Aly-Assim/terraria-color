"""Read the private wall list and resolve exact objects on the official wiki.

The ID indexes enrich names supplied by the workbook; they never define the
catalogue. Group-page ID ranges are deliberately not interpreted positionally.
"""
from collections import Counter
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import time
from urllib.parse import quote, unquote, urljoin, urlparse

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from PIL import Image
import requests

from scripts.paths import PROJECT_ROOT
# These helpers have no side effects; never call the historical main/setup.
from scripts.maintenance.import_catalog import normalize_name, safe_filename

WIKI = "https://terraria.wiki.gg"
WORKBOOK = PROJECT_ROOT / "source/all_items_terraria_145.xlsx"
USER_AGENT = "TerrariaColor/1.0 (wall catalogue importer; https://github.com/Aly-Assim/terraria-color)"


def read_walls(path: Path = WORKBOOK) -> tuple[list[dict], dict]:
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        sheets = [s for s in book.sheetnames if s.casefold() == "walls"]
        if len(sheets) != 1:
            raise ValueError(f"Expected one Walls sheet; found {book.sheetnames}")
        rows = list(book[sheets[0]].values)
        entries = []
        for row_number, row in enumerate(rows[1:], 2):
            for col, value in enumerate(row[:-1]):
                if not isinstance(value, str) or not isinstance(row[col + 1], str):
                    continue
                match = re.fullmatch(r'=(?:_xlfn\.)?IMAGE\("([^"]+)"\)', row[col + 1].strip(), re.I)
                if not match:
                    continue
                headers = [v for v in rows[0][:col + 1] if isinstance(v, str) and v.strip()]
                entries.append(dict(source_name=value.strip(), category_name=headers[-1] if headers else None,
                                    excel_image_url=match[1], source_row=row_number, source_column=col + 1))
        counts = Counter(normalize_name(e["source_name"]) for e in entries)
        analysis = dict(sheets=book.sheetnames, sheet=sheets[0], rows=len(rows),
                        columns=max(map(len, rows)), entries=len(entries),
                        categories=dict(Counter(e["category_name"] for e in entries)),
                        duplicate_names=[n for n, count in counts.items() if count > 1])
    finally:
        book.close()
    if not entries:
        raise ValueError("No name + IMAGE formula pairs in Walls")
    if analysis["duplicate_names"]:
        raise ValueError(f"Duplicate source names require review: {analysis['duplicate_names']}")
    return entries, analysis


def official_url(url: str) -> str:
    url = urljoin(WIKI, url)
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.netloc != "terraria.wiki.gg":
        raise ValueError(f"Only HTTPS English official wiki URLs are accepted: {url}")
    return url


def media_filename(url: str) -> str:
    parts = unquote(urlparse(url).path).split("/")
    return parts[-2] if "thumb" in parts else parts[-1]


def original_media(url: str) -> str:
    # Use the original filename, never a scaled thumbnail. The current wiki
    # exposes originals directly under /images/ (old hashed URLs redirect).
    return official_url("/images/" + quote(media_filename(url).replace(" ", "_"), safe="_") + "?format=original")


def image_bytes(data: bytes) -> str:
    with Image.open(BytesIO(data)) as image:
        extension = {"PNG": ".png", "GIF": ".gif", "JPEG": ".jpg", "WEBP": ".webp"}.get(image.format)
        if not extension or not all(image.size):
            raise ValueError("Unsupported/empty image")
        image.verify()
    with Image.open(BytesIO(data)) as image:
        for frame in range(getattr(image, "n_frames", 1)):
            image.seek(frame)
            image.load()
    return extension


class WikiClient:
    def __init__(self, cache: Path | None = None, refresh: bool = False, delay: float = 0.5):
        self.cache = cache or PROJECT_ROOT / "cache/walls/http"
        self.refresh = refresh
        self.delay = delay
        self.last_request = 0.0
        self.memory: dict[str, bytes] = {}
        self.session = requests.Session()
        self.session.headers["User-Agent"] = USER_AGENT

    def get(self, url: str) -> bytes:
        url = official_url(url).split("#", 1)[0]
        if url in self.memory:
            return self.memory[url]
        path = self.cache / sha256(url.encode()).hexdigest()
        if path.is_file() and not self.refresh:
            data = path.read_bytes()
        else:
            time.sleep(max(0, self.delay - (time.monotonic() - self.last_request)))
            response = self.session.get(url, timeout=(15, 60))
            self.last_request = time.monotonic()
            response.raise_for_status()
            official_url(response.url)
            data = response.content
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(data)
        self.memory[url] = data
        return data

    def page(self, url: str):
        soup = BeautifulSoup(self.get(url), "html.parser")
        if not soup.select_one(".mw-parser-output"):
            raise ValueError(f"No wiki article content: {url}")
        return soup


def parse_id_indexes(walls_html: bytes, items_html: bytes) -> tuple[dict, dict]:
    walls, items = {}, {}
    for tr in BeautifulSoup(walls_html, "html.parser").select("#table-walls tr"):
        cells = tr.find_all("td", recursive=False)
        if len(cells) < 5 or not cells[0].get_text(strip=True).isdigit():
            continue
        link = cells[1].find("a")
        img = cells[2].find("img")
        if not link:
            continue
        name = link.get_text(strip=True)
        wall = dict(wall_id=int(cells[0].get_text(strip=True)),
                    is_safe=1 if cells[4].select_one(".t-yes") else 0 if cells[4].select_one(".t-no") else None,
                    internal_name=cells[3].get_text(strip=True), name=name,
                    page_url=official_url(link["href"]),
                    world_image_url=original_media(img["src"]) if img else None)
        walls.setdefault(normalize_name(name), []).append(wall)
    for tr in BeautifulSoup(items_html, "html.parser").select("table.terraria tr"):
        cells = tr.find_all("td", recursive=False)
        if len(cells) < 3 or not cells[0].get_text(strip=True).isdigit():
            continue
        link = cells[1].find("a")
        if link:
            name = link.get_text(strip=True)
            items.setdefault(normalize_name(name), []).append(dict(
                item_id=int(cells[0].get_text(strip=True)), name=name, page_url=official_url(link["href"])))
    if not walls or not items:
        raise ValueError("Wiki ID index layout changed; refusing unverified metadata")
    return walls, items


class WallResolver:
    def __init__(self, client: WikiClient):
        self.client = client
        self.walls, self.items = parse_id_indexes(client.get(WIKI + "/wiki/Wall_IDs"),
                                                 client.get(WIKI + "/wiki/Item_IDs"))
        self.by_id = {w["wall_id"]: w for values in self.walls.values() for w in values}

    def resolve(self, entry: dict) -> dict:
        row = dict(entry, name=entry["source_name"], internal_item_id=None, internal_wall_id=None,
                   inventory_image_url=None, world_image_url=None, wall_ids=[], status="ok", problem=None,
                   grouped_page=False)
        problems = []
        key = normalize_name(row["name"])
        # Source typos such as "Spooky Wood" are retained in source_name. Only
        # accept the workbook's placed filename as an alias if Wall IDs confirms it.
        if key not in self.walls and entry["category_name"] != "Unsafe Walls":
            alias = normalize_name(media_filename(entry["excel_image_url"]))
            if alias in self.walls and alias != key:
                row["name"] = self.walls[alias][0]["name"]
                problems.append(f"source_name_alias: {entry['source_name']} -> {row['name']}")
                key = alias
        item_matches = self.items.get(key, [])
        if len(item_matches) > 1:
            raise ValueError(f"Ambiguous Item ID for {row['name']}")
        wall_matches = self.walls.get(key, [])
        url = item_matches[0]["page_url"] if item_matches else (
            wall_matches[0]["page_url"] if wall_matches else WIKI + "/wiki/" + quote(row["name"].replace(" ", "_")))
        url = entry.get("page_url") or url
        soup = self.client.page(url)
        canonical = soup.select_one('link[rel="canonical"]')
        row["page_url"] = official_url(canonical["href"] if canonical else url)
        exact_boxes = [box for box in soup.select(".infobox")
                       if box.select_one(".title") and normalize_name(box.select_one(".title").get_text(" ", strip=True)) == key]
        if len(exact_boxes) > 1:
            raise ValueError(f"Multiple matching infoboxes: {row['name']}")
        box = exact_boxes[0] if exact_boxes else None
        row["grouped_page"] = box is None
        if box is not None:
            # Variant-specific infobox (e.g. Natural Dirt Wall). Never parse
            # ranges or the unrelated first infobox on a group page.
            for line in box.select(".ids li"):
                if "Wall ID" not in line.get_text(" ", strip=True):
                    continue
                ids = line.find("b")
                if ids is not None and not re.search(r"\d\s*[–-]\s*\d", ids.get_text()):
                    values = [int(v) for v in re.findall(r"\d+", ids.get_text())]
                    if values and all(v in self.by_id for v in values):
                        wall_matches = [self.by_id[v] for v in values]
        if not wall_matches and entry["category_name"] == "Unsafe Walls" and item_matches:
            # Cursed Dungeon variants: the exact Shimmer output row establishes
            # the base wall; the workbook supplies the same placed texture.
            # Select only that base wall's unsafe ID, never a sibling's range.
            placed_key = normalize_name(media_filename(entry["excel_image_url"]))
            for result in soup.select("td.result[data-sort-value]"):
                if normalize_name(result["data-sort-value"]) != key:
                    continue
                ingredients = result.parent.select("td.ingredients img[src]")
                if any(normalize_name(media_filename(img["src"])) == placed_key for img in ingredients):
                    candidates = [w for w in self.walls.get(placed_key, []) if w["is_safe"] == 0]
                    if len(candidates) == 1:
                        wall_matches = candidates
        if not wall_matches:
            problems.append("missing_wall_id")
        row["wall_ids"] = wall_matches
        preferred = [w for w in wall_matches if w["is_safe"] == 1]
        chosen = preferred if preferred else wall_matches
        if len(chosen) == 1:
            row["internal_wall_id"] = chosen[0]["wall_id"]
        elif chosen:
            problems.append("ambiguous_main_wall_id")
        if item_matches:
            row["internal_item_id"] = item_matches[0]["item_id"]
        elif box is not None:
            match = re.search(r"Internal Item ID\s*:\s*(\d+)(?![\d–-])(?:\s|$)", box.get_text(" ", strip=True))
            if match:
                row["internal_item_id"] = int(match[1])
        images = (box or soup.select_one(".mw-parser-output")).select("img[src]")
        inv, world = set(), set()
        for image in images:
            filename = media_filename(image["src"])
            if normalize_name(filename) == key:
                (world if "(placed)" in filename.lower() else inv).add(original_media(image["src"]))
        if len(inv) == 1 and row["internal_item_id"] is not None:
            row["inventory_image_url"] = inv.pop()
        elif len(inv) > 1:
            problems.append("ambiguous_inventory_media")
        if len(world) == 1:
            row["world_image_url"] = world.pop()
        elif len(world) > 1:
            problems.append("ambiguous_world_media")
        elif len(chosen) == 1:
            # An exact ID row also establishes its placed-image identity.
            row["world_image_url"] = chosen[0]["world_image_url"]
        if row["world_image_url"] is None:
            problems.append("missing_world_media")
        if row["inventory_image_url"] is None:
            problems.append("inventory_item_not_found" if row["internal_item_id"] is None and wall_matches else "missing_inventory_media")
        row["color_image_url"] = row["world_image_url"]
        row["problem"] = "; ".join(problems) or None
        row["status"] = "review" if problems else "ok"
        return row
